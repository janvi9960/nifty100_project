from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# Allow importing from project root
PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from src.screener.engine import ScreenerEngine


OUTPUT_DIR = PROJECT_ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "screener_output.xlsx"


PRESETS = [
    "Quality Compounder",
    "Value Pick",
    "Growth Accelerator",
    "Dividend Champion",
    "Debt-Free Blue Chip",
    "Turnaround Watch",
]


# Columns required in the final Excel workbook
KPI_COLUMNS = [
    "company_id",
    "company_name",
    "broad_sector",
    "sub_sector",
    "return_on_equity",
    "return_on_capital",
    "return_on_assets",
    "net_profit_margin",
    "operating_profit_margin",
    "debt_to_equity",
    "interest_coverage",
    "free_cash_flow",
    "cash_from_operations_cr",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "asset_turnover",
    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",
    "composite_quality_score",
]


def safe_sheet_name(name):
    """Convert preset name into a valid Excel sheet name."""
    return name[:31]


def export_screener():
    print("=" * 60)
    print("NIFTY 100 SCREENER - EXCEL EXPORT")
    print("=" * 60)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    engine = ScreenerEngine()

    exported = {}

    for preset in PRESETS:
        print(f"\nRunning: {preset}")

        df = engine.run_preset(preset)

        if df.empty:
            print(f"WARNING: {preset} returned 0 companies.")
            continue

        # Keep only columns that actually exist
        columns = [col for col in KPI_COLUMNS if col in df.columns]
        export_df = df[columns].copy()

        # Sort by composite score
        if "composite_quality_score" in export_df.columns:
            export_df = export_df.sort_values(
                "composite_quality_score",
                ascending=False,
                na_position="last",
            )

        exported[preset] = export_df

        print(f"Companies: {len(export_df)}")

    if not exported:
        raise RuntimeError("No screener results were generated.")

    # Write all sheets
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for preset, df in exported.items():
            sheet_name = safe_sheet_name(preset)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    # ---------------------------------------------------------
    # Excel formatting
    # ---------------------------------------------------------
    workbook = load_workbook(OUTPUT_FILE)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    pass_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    fail_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    for worksheet in workbook.worksheets:

        # Header formatting
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Find composite score column
        score_column = None

        for cell in worksheet[1]:
            if cell.value == "composite_quality_score":
                score_column = cell.column
                break

        # Green/red formatting based on composite score
        if score_column is not None:
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=score_column)

                if isinstance(cell.value, (int, float)):
                    if cell.value >= 50:
                        cell.fill = pass_fill
                    else:
                        cell.fill = fail_fill

        # Number formatting
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

        # Adjust column widths
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                35,
            )

    workbook.save(OUTPUT_FILE)

    print("\n" + "=" * 60)
    print("EXCEL EXPORT COMPLETED")
    print("=" * 60)
    print(f"File created: {OUTPUT_FILE}")
    print(f"Sheets created: {len(exported)}")

    for preset, df in exported.items():
        print(f"  {preset}: {len(df)} companies")

    print("=" * 60)


if __name__ == "__main__":
    export_screener()
    